# -*- coding: utf-8 -*-
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text, title="ĐIỂM NHẤN CỐT LÕI", hex_bg="F0F4F8", border_color="1F4E79"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(10.5)
    r_text.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="6C757D"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E9ECEF"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E9ECEF"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E9ECEF"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x21, 0x25, 0x29)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = "Calibri"
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(15)
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return h

def add_para(doc, text, bold_prefix=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(11)
        r_b.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    r_t = p.add_run(text)
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(11)
    r_t.italic = italic
    r_t.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_bullet(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(10.5)
        r_b.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    r_t = p.add_run(text)
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(10.5)
    r_t.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def create_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)

    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def main():
    doc = Document()

    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("BÁO CÁO PHƯƠNG PHÁP & KIẾN TRÚC HỆ THỐNG")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(19)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("CONFIGURABLE LECTURE SCRIPT & VISUAL INTENT GENERATION (CLSG-IR)\nCầu Nối Biểu Diễn Trung Gian (IR) Giữa Tài Liệu Học Tập và Hệ Thống AI Video Generation")
    r_sub.bold = True
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12.5)
    r_sub.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(20)
    r_meta = p_meta.add_run("Vai trò: Instructional Designer & Educational Content Architect | Đối sánh: EduCraft (CIKM 2025)")
    r_meta.italic = True
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = RGBColor(0xBD, 0xD7, 0xEE)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1
    format_heading(doc, "1. TỔNG QUAN VÀ ĐỘNG LỰC NGHIÊN CỨU (INTRODUCTION & MOTIVATION)", 1)
    add_para(doc, 
        "Trong làn sóng ứng dụng AI tạo sinh vào giáo dục, nhu cầu chuyển đổi tự động các tài liệu học tập thô (PPTX, DOCX, Markdown, PDF) "
        "thành bài giảng đa phương tiện đang gia tăng mạnh mẽ. Tuy nhiên, các hệ thống hiện nay đang tồn tại sự đứt gãy nghiêm trọng:"
    )

    add_bullet(doc, "EduCraft cố gắng dùng VLM để đọc từng slide bài giảng. Cách tiếp cận này bị trói cứng vào bố cục slide (1 slide = 1 đoạn kịch bản rời rạc), tiêu tốn tài nguyên GPU khổng lồ, thường xuyên bị ảo giác thị giác và không thể co giãn bài giảng theo thời lượng hay đối tượng người học.", bold_prefix="1. Nhược điểm của EduCraft (CIKM 2025): ")
    add_bullet(doc, "Các công cụ tạo video hiện đại (Sora, Runway, Midjourney, HeyGen) rất xuất sắc về chất lượng dựng nhưng 'mù hoàn toàn về mặt sư phạm' (Pedagogically Blind). Chúng không thể tự nhận biết khi nào học viên gặp khó khăn nhận thức để chèn biểu đồ so sánh hay sơ đồ quy trình.", bold_prefix="2. Sự đứt gãy với AI Video Generation: ")

    add_callout(doc, 
        "Hệ thống CLSG-IR đóng vai trò là Lớp Biểu Diễn Trung Gian (Intermediate Representation - IR). "
        "Mục tiêu KHÔNG phải tạo storyboard hoàn chỉnh (không tạo camera direction, scene transition hay shot list). "
        "Hệ thống tạo ra một hợp đồng dữ liệu gồm Lời thoại giảng giải chuẩn văn phong nói (Narration) kết hợp với Siêu dữ liệu ý định trực quan sư phạm (Visual Intent Metadata).",
        title="ĐỊNH VỊ NỀN TẢNG (CORE PARADIGM)"
    )

    # Section 2
    format_heading(doc, "2. NGUYÊN TẮC THIẾT KẾ VÀ HỆ THỐNG PHÂN LOẠI TRỰC QUAN (TAXONOMY & RULES)", 1)
    add_para(doc, "Hệ thống vận hành theo 5 nguyên tắc sư phạm bắt buộc:")
    add_bullet(doc, "Nội dung và mục tiêu học tập (Learning Goals) là ưu tiên cao nhất.", bold_prefix="Nguyên tắc 1: ")
    add_bullet(doc, "Chỉ đề xuất hình ảnh khi nó thực sự giúp người học hiểu bài tốt hơn (giảm tải nhận thức).", bold_prefix="Nguyên tắc 2: ")
    add_bullet(doc, "Tuyệt đối KHÔNG tạo hình ảnh chỉ nhằm mục đích trang trí hình thức.", bold_prefix="Nguyên tắc 3: ")
    add_bullet(doc, "Mỗi visual cue phải trả lời rõ: Học viên cần hiểu gì? Điểm nào khó hình dung? Loại visual nào hỗ trợ tốt nhất?", bold_prefix="Nguyên tắc 4: ")
    add_bullet(doc, "Phân định trách nhiệm tuyệt đối: Không sinh chỉ thị camera, góc quay, hiệu ứng chuyển cảnh hay hướng dẫn dựng video.", bold_prefix="Nguyên tắc 5: ")

    format_heading(doc, "2.1. Bảng phân loại 13 loại trực quan chuẩn (Visual Type Taxonomy)", 2)
    add_para(doc, "Mọi visual cue do hệ thống sinh ra bắt buộc phải thuộc một trong 13 phân loại sau:")

    headers_tax = ["Visual Type", "Mục đích Sư phạm Chính", "Quy tắc Kích hoạt (Selection Rules)"]
    data_tax = [
        ["diagram", "Trực quan hóa cấu trúc, bộ phận của một thực thể tĩnh.", "Định nghĩa khái niệm, kiến trúc hệ thống tĩnh."],
        ["flowchart", "Mô tả chuỗi các bước thực hiện tuần tự hoặc rẽ nhánh logic.", "Quy trình nhiều bước, giải thuật, cây quyết định."],
        ["comparison", "Đặt cạnh nhau để làm nổi bật sự khác biệt bản chất.", "So sánh 2 hoặc nhiều khái niệm (vd: Supervised vs Unsupervised)."],
        ["infographic", "Tổng hợp toàn cảnh tri thức đa chiều kết hợp số liệu.", "Bức tranh tổng quan hệ sinh thái, nguyên lý phức hợp."],
        ["chart", "Trực quan hóa các dữ liệu định lượng, tỷ lệ và xu hướng.", "Số liệu thống kê, biểu đồ biến thiên, phân tích hiệu năng."],
        ["table", "Cấu trúc hóa thông tin đa thuộc tính dạng hàng và cột.", "Đối chiếu thông số kỹ thuật, bảng dữ liệu mẫu."],
        ["screenshot", "Minh chứng giao diện thực tế của phần mềm hoặc hệ điều hành.", "Thao tác trên phần mềm, màn hình lập trình, lệnh console."],
        ["illustration", "Tranh vẽ ẩn dụ trực quan hóa ý niệm trừu tượng.", "Ẩn dụ đời sống giải thích cơ chế kỹ thuật phức tạp."],
        ["real_world_example", "Hình ảnh vật thể, hiện trường hoặc tình huống thực tế.", "Ứng dụng thực tiễn ngoài đời sống, sản phẩm công nghiệp."],
        ["timeline", "Biểu diễn diễn biến sự kiện theo trục thời gian tuyến tính.", "Lịch sử phát triển công nghệ, các mốc tiến hóa."],
        ["equation", "Hiển thị công thức toán học, ký hiệu logic chặt chẽ.", "Định lý, phương trình, hàm mất mát (loss function)."],
        ["process_visualization", "Mô phỏng động lực dòng chảy hoạt động của hệ thống.", "Cơ chế truyền dữ liệu qua các tầng mạng, chu trình động cơ."],
        ["concept_map", "Bản đồ mạng lưới các khái niệm và mối liên hệ ngữ nghĩa.", "Tổng kết mạng lưới tri thức liên ngành, cây phân loại."]
    ]
    create_table(doc, headers_tax, data_tax, col_widths=[1.6, 2.7, 2.2])

    # Section 3
    format_heading(doc, "3. KIẾN TRÚC PIPELINE HOÀN CHỈNH (SYSTEM PIPELINE & ARCHITECTURE)", 1)
    add_para(doc, "Kiến trúc hệ thống gồm 4 tầng phân tách độc lập, khép kín với vòng lặp kiểm định chất lượng:")

    # Embedded Image
    try:
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(r"c:\Users\KIM THIEN\Desktop\T032\clsg_pipeline_architecture.png", width=Inches(6.5))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Hình 1: Sơ đồ Pipeline tổng thể hệ thống CLSG-IR làm cầu nối cho AI Video Generation")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.italic = True
        r_cap.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    except Exception as e:
        print(f"Lỗi chèn ảnh: {e}")

    add_bullet(doc, "Parser theo luật thuần túy (AST) trích xuất toàn bộ cấu trúc Heading, Paragraph, Table từ PPTX, DOCX, Markdown. Không dùng VLM, độ trễ < 100ms, chi phí 0 USD.", bold_prefix="Module 1 (Content Extractor): ")
    add_bullet(doc, "Tính toán ngân sách số từ W_target = Duration × 140 WPM; chia bài giảng thành các section cân đối; xác định learning_goal cụ thể theo Bloom's Taxonomy.", bold_prefix="Module 2 (Instructional Planner): ")
    add_bullet(doc, "Sinh lời giảng hoàn chỉnh (narration) văn phong nói tự nhiên, đồng thời áp dụng Visual Selection Rules để sinh mảng visual_cues tương ứng (chỉ khi có nhu cầu nhận thức thực sự).", bold_prefix="Module 3 (Script & Visual Generator): ")
    add_bullet(doc, "Đo sai số độ dài lời thoại (|ΔW|/W_target <= 15%), kiểm tra tính chuẩn xác khoa học, loại bỏ các visual mang tính trang trí thừa thãi.", bold_prefix="Module 4 (Quality & Visual Guard): ")

    # Section 4
    format_heading(doc, "4. ĐẶC TẢ SCHEMA DỮ LIỆU ĐẦU RA (OUTPUT JSON SCHEMA)", 1)
    add_para(doc, "Mỗi section bài giảng được đóng gói thành một đối tượng JSON chuẩn mực:")

    schema_code = (
        '{\n'
        '  "section_id": "sec_01",\n'
        '  "topic": "Phân biệt Học có giám sát và Học không giám sát",\n'
        '  "learning_goal": "Người học phân biệt được bản chất khác nhau giữa 2 phương pháp dựa trên sự hiện diện của nhãn dữ liệu.",\n'
        '  "narration": "Chào các bạn. Khi bước chân vào thế giới Học máy, câu hỏi đầu tiên luôn là: Máy tính sẽ học từ loại dữ liệu nào? Hãy hình dung bạn đang dạy một đứa trẻ nhận biết trái cây...",\n'
        '  "visual_cues": [\n'
        '    {\n'
        '      "visual_need": true,\n'
        '      "visual_type": "comparison",\n'
        '      "visual_purpose": "Làm nổi bật sự khác biệt bản chất giữa dữ liệu có nhãn và không nhãn.",\n'
        '      "content_focus": "Hai luồng dữ liệu song song: Luồng 1 (Dữ liệu + Nhãn -> Mô hình dự đoán); Luồng 2 (Dữ liệu thô -> Thuật toán tự gom cụm).",\n'
        '      "learning_support": "Giúp người học khắc phục điểm khó hình dung nhất về cách dữ liệu đi vào mô hình ở hai bài toán.",\n'
        '      "importance_level": "high"\n'
        '    },\n'
        '    {\n'
        '      "visual_need": true,\n'
        '      "visual_type": "real_world_example",\n'
        '      "visual_purpose": "Neo giữ khái niệm trừu tượng vào trải nghiệm đời thực quen thuộc.",\n'
        '      "content_focus": "Một giỏ trái cây dán nhãn tên đối chiếu với một giỏ trái cây tự động phân tách theo màu sắc và hình dáng.",\n'
        '      "learning_support": "Tạo cầu nối trực quan giúp người học không chuyên ghi nhớ bản chất nhanh chóng.",\n'
        '      "importance_level": "medium"\n'
        '    }\n'
        '  ],\n'
        '  "section_summary": "Học có giám sát dùng dữ liệu có nhãn để dự đoán; Học không giám sát dùng dữ liệu không nhãn để tự tìm quy luật."\n'
        '}'
    )
    add_code_block(doc, schema_code)

    # Section 5
    format_heading(doc, "5. BẢNG ĐỐI TRỌNG VỚI EDU CRAFT (CIKM 2025)", 1)
    add_para(doc, "Bảng đối sánh làm nổi bật bước tiến đột phá của phương pháp đề xuất:")

    headers_comp = ["Tiêu chí So sánh", "EduCraft (CIKM 2025)", "CLSG-IR (Phương pháp Đề xuất)"]
    data_comp = [
        ["Mục tiêu cốt lõi", "Sinh kịch bản đọc theo từng trang slide", "Sinh Biểu diễn Trung gian (IR) cho Video AI"],
        ["Xử lý thị giác", "Cồng kềnh (VLM, OCR slide, Bounding box)", "Zero-Vision Extractor (Nhanh, không tốn GPU)"],
        ["Siêu dữ liệu hình ảnh", "Không có hoặc mô tả slide tĩnh có sẵn", "Visual Intent Metadata (13 Taxonomies + Rules)"],
        ["Tính năng sư phạm", "Bị ràng buộc bởi nội dung người vẽ slide", "Chủ động đề xuất visual hỗ trợ nhận thức"],
        ["Kiểm soát thời lượng", "Cố định theo số lượng slide", "Co giãn chính xác theo ngân sách từ (WPM)"],
        ["Khả năng tích hợp", "Khó kết nối các công cụ sinh video", "Nối ghép trực tiếp với Midjourney, Sora, Runway"]
    ]
    create_table(doc, headers_comp, data_comp, col_widths=[1.8, 2.3, 2.4])

    # Section 6
    format_heading(doc, "6. HỆ THỐNG ĐÁNH GIÁ THỰC NGHIỆM (EVALUATION METRICS)", 1)
    add_para(doc, "Hệ thống được đánh giá dựa trên 4 chỉ số khoa học định lượng:")
    add_bullet(doc, "Chuyên gia giáo dục hoặc Critic LLM đánh giá tỷ lệ các visual cues đề xuất thực sự giải tỏa điểm nghẽn nhận thức, không mang tính trang trí vô bổ.", bold_prefix="1. Visual Necessity Score (VNS): ")
    add_bullet(doc, "Tỷ lệ 100% các visual_type được gán đúng theo 13 loại trong Taxonomy quy định.", bold_prefix="2. Taxonomy Compliance Rate (TCR): ")
    add_bullet(doc, "Độ khớp về ngữ nghĩa giữa content_focus trong visual cue với đoạn narration đang được phát biểu.", bold_prefix="3. Narration-Visual Alignment: ")
    add_bullet(doc, "DAR = 1 - (|W_actual - W_target| / W_target). Đảm bảo thời lượng lời giảng tuân thủ mục tiêu của người dùng.", bold_prefix="4. Duration Adherence Rate (DAR): ")

    # Section 7
    format_heading(doc, "7. KẾT LUẬN", 1)
    add_para(doc, 
        "Hệ thống CLSG-IR tái định hình bài toán Lecture Script Generation thành bài toán thiết kế Intermediate Representation cho Video AI. "
        "Bằng cách kết hợp giữa lời giảng giàu tính đối thoại với siêu dữ liệu trực quan sư phạm nghiêm ngặt, khung phương pháp này loại bỏ hoàn toàn các gánh nặng "
        "của EduCraft, mở đường cho việc tự động hóa sản xuất video giáo dục chất lượng cao ở quy mô lớn."
    )

    output_filename = r"c:\Users\KIM THIEN\Desktop\T032\Bao_Cao_Phuong_Phap_De_Xuat_CLSG.docx"
    doc.save(output_filename)
    print(f"Document successfully created at: {output_filename}")

if __name__ == "__main__":
    main()
