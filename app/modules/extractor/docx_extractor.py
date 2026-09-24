# app/modules/extractor/docx_extractor.py
"""
Deterministic DOCX Extractor for CLSG-IR Module 1.
Extracts document headings, paragraphs, bullet points, and tables.
Implements Multimodal Document Representation:
- DOCUMENT IR = SOURCE OF TRUTH (DocumentElement, StructuredTable, SlideIR)
- MARKDOWN = LLM-FRIENDLY VIEW
- SEMANTIC CHUNK = UNIT OF REASONING
"""
import io
import time
import uuid
from typing import Union, List
from pathlib import Path
from docx import Document

from app.models.document import CanonicalDocumentTree, DocumentSection, ContentElement, VisualElement
from app.models.document_ir import DocumentIR, SlideIR, DocumentElement, StructuredTable
from app.modules.extractor.base import BaseExtractor
from app.modules.extractor.element_classifier import ElementClassifier
from app.modules.extractor.relationship_builder import RelationshipBuilder
from app.modules.extractor.normalizer import ContentNormalizer
from app.modules.extractor.chunker import SemanticChunker

class DOCXExtractor(BaseExtractor):
    def __init__(self):
        self.classifier = ElementClassifier()
        self.rel_builder = RelationshipBuilder()
        self.normalizer = ContentNormalizer()
        self.chunker = SemanticChunker()

    def extract(self, source: Union[str, Path, bytes], filename: str) -> CanonicalDocumentTree:
        start_time = time.perf_counter()
        
        if isinstance(source, bytes):
            doc = Document(io.BytesIO(source))
        else:
            doc = Document(str(source))
            
        doc_id = f"doc_{uuid.uuid4().hex[:8]}"
        overall_title = Path(filename).stem.replace("_", " ").title()

        slides_ir: List[SlideIR] = []
        legacy_sections: List[DocumentSection] = []

        current_sec_id = "S1"
        current_title = "Introduction"
        current_elements: List[DocumentElement] = []
        legacy_elements: List[ContentElement] = []
        raw_texts: List[str] = []
        el_idx = 1
        sec_idx = 1

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""

            if "heading 1" in style_name or "heading 2" in style_name:
                if current_elements:
                    relationships = self.rel_builder.build_slide_relationships(current_elements)
                    slides_ir.append(SlideIR(
                        slide_id=sec_idx,
                        section_id=current_sec_id,
                        title=current_title,
                        elements=current_elements,
                        relationships=relationships
                    ))
                    legacy_sections.append(DocumentSection(
                        section_id=current_sec_id,
                        title=current_title,
                        order=sec_idx,
                        elements=legacy_elements,
                        raw_text="\n".join(raw_texts)
                    ))
                    sec_idx += 1
                    current_sec_id = f"S{sec_idx}"
                    current_elements = []
                    legacy_elements = []
                    raw_texts = []
                    el_idx = 1
                
                current_title = text

            el_type, el_role, conf = self.classifier.classify_text_element(text, is_title=("heading" in style_name))
            current_elements.append(DocumentElement(
                id=f"{doc_id}_S{sec_idx:02d}_el{el_idx:02d}",
                type=el_type,
                role=el_role,
                classification_confidence=conf,
                reading_order=el_idx,
                content=text,
                source_slide=sec_idx
            ))
            legacy_elements.append(ContentElement(
                element_id=f"{current_sec_id}_el_{el_idx:02d}",
                type="paragraph",
                text=text
            ))
            raw_texts.append(text)
            el_idx += 1

        # Process any tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])
            if table_rows:
                headers = table_rows[0]
                rows_data = table_rows[1:] if len(table_rows) > 1 else table_rows
                table_str = "\n".join(" | ".join(r) for r in table_rows)
                st_table = StructuredTable(
                    table_id=f"tbl_{doc_id}_S{sec_idx:02d}_{el_idx:02d}",
                    columns=headers,
                    rows=rows_data,
                    headers=headers
                )
                current_elements.append(DocumentElement(
                    id=f"{doc_id}_S{sec_idx:02d}_el{el_idx:02d}",
                    type="table",
                    role="comparison",
                    classification_confidence=0.95,
                    reading_order=el_idx,
                    content=table_str,
                    structured_table=st_table,
                    source_slide=sec_idx
                ))
                legacy_elements.append(ContentElement(
                    element_id=f"{current_sec_id}_el_{el_idx:02d}",
                    type="table",
                    text=table_str,
                    metadata={"rows": len(table.rows), "cols": len(table.columns)}
                ))
                raw_texts.append(table_str)
                el_idx += 1

        if current_elements:
            relationships = self.rel_builder.build_slide_relationships(current_elements)
            slides_ir.append(SlideIR(
                slide_id=sec_idx,
                section_id=current_sec_id,
                title=current_title,
                elements=current_elements,
                relationships=relationships
            ))
            legacy_sections.append(DocumentSection(
                section_id=current_sec_id,
                title=current_title,
                order=sec_idx,
                elements=legacy_elements,
                raw_text="\n".join(raw_texts)
            ))

        doc_ir = DocumentIR(
            document_id=doc_id,
            title=overall_title,
            source_type="docx",
            source_filename=filename,
            total_slides=len(slides_ir),
            slides=slides_ir
        )

        canonical_markdown = self.normalizer.render_from_document_ir(doc_ir)
        semantic_chunks = self.chunker.chunk_from_document_ir(doc_ir)

        elapsed_ms = (time.perf_counter() - start_time) * 1000.0

        return CanonicalDocumentTree(
            document_id=doc_id,
            title=overall_title,
            source_type="docx",
            source_filename=filename,
            total_sections=len(legacy_sections),
            sections=legacy_sections,
            extraction_time_ms=round(elapsed_ms, 2),
            canonical_markdown=canonical_markdown,
            semantic_chunks=semantic_chunks,
            visual_elements=[],
            document_ir=doc_ir
        )
